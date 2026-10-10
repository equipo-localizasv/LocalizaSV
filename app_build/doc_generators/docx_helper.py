import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

NAVY = RGBColor(10, 27, 57)        # #0A1B39
CYAN = RGBColor(2, 132, 199)       # #0284C7
SLATE = RGBColor(71, 85, 105)      # #475569
CHARCOAL = RGBColor(30, 41, 59)    # #1E293B
WHITE = RGBColor(255, 255, 255)
GREEN = RGBColor(16, 185, 129)

def set_cell_background(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_document(title, subtitle, doc_code="DOC-LOCSV-2026", version="2.5.0"):
    doc = Document()
    
    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run(f"LOCALIZASV — C4I EL SALVADOR | {doc_code}")
        hrun.font.name = 'Arial'
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = SLATE
        
        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        frun1 = fp.add_run("CONFIDENCIAL • SISTEMA NACIONAL DE BÚSQUEDA FORENSE SOLIDARIA")
        frun1.font.name = 'Arial'
        frun1.font.size = Pt(8)
        frun1.font.color.rgb = SLATE

    # Portada / Banner Superior
    p_meta = doc.add_paragraph()
    r_meta = p_meta.add_run("REPÚBLICA DE EL SALVADOR • SISTEMA C4I DE SEGURIDAD PÚBLICA")
    r_meta.font.name = 'Arial'
    r_meta.font.size = Pt(9)
    r_meta.font.bold = True
    r_meta.font.color.rgb = CYAN
    p_meta.paragraph_format.space_after = Pt(4)

    p_title = doc.add_paragraph()
    r_title = p_title.add_run(title)
    r_title.font.name = 'Arial'
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = NAVY
    p_title.paragraph_format.space_after = Pt(4)

    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run(subtitle)
    r_sub.font.name = 'Arial'
    r_sub.font.size = Pt(12)
    r_sub.font.color.rgb = SLATE
    p_sub.paragraph_format.space_after = Pt(14)

    # Metadata table (Código, Versión, Fecha, Clasificación)
    meta_table = doc.add_table(rows=2, cols=4)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_headers = ["Código de Documento", "Versión", "Fecha de Emisión", "Clasificación"]
    meta_values = [doc_code, f"v{version}", "Octubre 2026", "Oficial / Técnico"]

    for i in range(4):
        c_h = meta_table.cell(0, i)
        c_h.text = meta_headers[i]
        set_cell_background(c_h, "0F172A")
        set_cell_margins(c_h, 80, 80, 100, 100)
        c_h.paragraphs[0].runs[0].font.bold = True
        c_h.paragraphs[0].runs[0].font.size = Pt(8.5)
        c_h.paragraphs[0].runs[0].font.color.rgb = WHITE
        
        c_v = meta_table.cell(1, i)
        c_v.text = meta_values[i]
        set_cell_background(c_v, "F1F5F9")
        set_cell_margins(c_v, 80, 80, 100, 100)
        c_v.paragraphs[0].runs[0].font.size = Pt(9)
        c_v.paragraphs[0].runs[0].font.color.rgb = CHARCOAL

    doc.add_paragraph().paragraph_format.space_after = Pt(16)
    return doc

def add_h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = NAVY
    return p

def add_h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(12.5)
    r.font.bold = True
    r.font.color.rgb = CYAN
    return p

def add_h3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = CHARCOAL
    return p

def add_p(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(10)
    r.font.color.rgb = CHARCOAL
    return p

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Arial'
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = CHARCOAL
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(10)
    r.font.color.rgb = CHARCOAL
    return p

def add_callout(doc, text, title="NOTA INFORMATIVA", alert_type="info"):
    colors = {
        "info": ("E0F2FE", "0284C7"),      # Cyan
        "warning": ("FEF3C7", "D97706"),   # Amber
        "success": ("DCFCE7", "16A34A"),   # Green
        "security": ("FEE2E2", "DC2626")   # Red
    }
    bg_hex, border_hex = colors.get(alert_type, colors["info"])
    
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, 120, 120, 180, 180)
    
    # Left border only
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="36" w:space="0" w:color="{border_hex}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r_t = p.add_run(f"📌 {title}: ")
    r_t.font.name = 'Arial'
    r_t.font.size = Pt(9.5)
    r_t.font.bold = True
    
    r_body = p.add_run(text)
    r_body.font.name = 'Arial'
    r_body.font.size = Pt(9.5)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_table(doc, headers, rows_data, col_widths=None):
    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Headers
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, 100, 100, 120, 120)
        p = cell.paragraphs[0]
        p.runs[0].font.name = 'Arial'
        p.runs[0].font.size = Pt(9)
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = WHITE
    
    # Data Rows
    for r_idx, row in enumerate(rows_data):
        bg = "FFFFFF" if r_idx % 2 == 0 else "F8FAFC"
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx + 1, c_idx)
            cell.text = str(val)
            set_cell_background(cell, bg)
            set_cell_margins(cell, 80, 80, 120, 120)
            p = cell.paragraphs[0]
            p.runs[0].font.name = 'Arial'
            p.runs[0].font.size = Pt(8.5)
            p.runs[0].font.color.rgb = CHARCOAL
            
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    return table

def add_image_if_exists(doc, image_path, caption="", width_inches=6.0):
    if os.path.exists(image_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run = p_img.add_run()
        run.add_picture(image_path, width=Inches(width_inches))
        
        if caption:
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(12)
            r_cap = p_cap.add_run(caption)
            r_cap.font.name = 'Arial'
            r_cap.font.size = Pt(8.5)
            r_cap.font.italic = True
            r_cap.font.color.rgb = SLATE
        return True
    return False
